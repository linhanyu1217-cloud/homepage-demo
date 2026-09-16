# -*- coding: utf-8 -*-
"""把 05_输出\\<周次>\\ 下的 md 交付物渲染成 Word（.docx）。

对应交付物
----------
    本周公司公告汇总.md  ->  本周公司公告周报.docx
    本周重要资讯.md      ->  本周重要资讯.docx

**Word 只有一份**：不再区分「正式版 / 备选版」，不再生成 本周公司公告备选.docx。
本周值得保留的公司公告全部整合进同一份《本周公司公告周报.docx》；
《本周公司公告汇总.md》是它的中间底稿（可 diff、可追溯），两者内容一一对应。

设计要点
--------
1. **模板原件只读**：先 shutil.copyfile 复制模板，再打开**副本**修改；
   模板自身一个字节都不动（脚本结束会校验模板 sha256 未变）。
2. **格式全部取自模板本体**：不新建样式，而是把模板里已有的段落
   （日标题段 / 【公司新闻】段 / List Paragraph 条目段）deepcopy 出来当"模具"，
   只替换文字。因此字体（Cambria/宋体、Calibri/宋体）、字号（16pt / 11pt）、
   行距（line 416 auto、keepNext/keepLines）、
   编号（decimal "%1."、ind left=440 hanging=440）与模板完全一致。
3. **编号逐组重启**：模板每天用**不同的 numId**——作者正是用「新列表实例」
   让编号每天从 1 重新开始。本脚本沿用该做法：每个日期分组动态克隆一份
   abstractNum 并注册新 numId（同时改写 nsid/tmpl，避免 Word 把两份定义
   当成同一个列表模板而继续编号）。
4. **清空正文**：写入前删除 body 内除 w:sectPr 以外的全部段落——模板里的旧
   公告内容即在此清掉；页面尺寸、页边距、页眉页脚不受影响。
5. **内容零改写**：所有文字直接来自 md，只去掉 markdown 标记
   （`#` / `-` / `>` / `|` / `**`），不改写、不新增、不删除任何数据。
   `**...**` 按 md 原意渲染为**加粗**（汇总条目本身不带 `**`，
   因此渲染结果与模板条目完全同款）。
6. **空行靠间距，不靠空段落**（2026-09-16 起）：全文**不插入任何空段落**，
   "一行空行"由段前间距 GAP_TWIPS=260 twips（13pt ≈ 11pt 正文的一行高）实现。
   模板每个日标题本来就带 `w:spacing w:before="260"`，所以 260 也正是模板
   自己的分节日距——所有层级由此做到**统一一行空行**：
       日标题→【公司新闻】/【公司新闻】→第 1 条/条目与条目之间/日块与日块之间
   全部 = 260 twips，不会出现连续两行以上空白。
   注意 `List Paragraph` 样式自带 `<w:contextualSpacing/>`，会**吞掉相邻同款
   段落之间的间距**；条目段落必须显式写 `<w:contextualSpacing w:val="0"/>`
   覆盖，否则公告会全部粘在一起。
7. **每个日期块另起一页**（2026-09-16 起）：日期块标题带 `<w:pageBreakBefore/>`
   （段落属性「段前分页」），因此新日期一定从新的一页顶部开始，
   永远不会落在上一页末尾；**第一个日期块不加**（第一页本来就在页首）。
   同一天的公告连续排列，不逐条分页；某天公告多时自然跨页，**不重复日期标题**。
   用 `pageBreakBefore` 而不是手工分页符 `<w:br w:type="page"/>`，是因为
   Word 在本段已位于页首时会忽略前者，**不会凭空多出一张空白页**。
   分页**不用空行凑**——全文仍然一个空段落都没有，空行只由 `w:spacing` 产生。

命令行
------
    python 脚本\\生成Word.py --week 2026-W37
    python 脚本\\生成Word.py --outdir 05_输出\\2026-W37 --force
    python 脚本\\生成Word.py --week 2026-W37 --force --only 汇总   # 只重做 Word 周报

作为模块导入
-----------
    from 生成Word import generate
    generate(out_dir=r"...\\05_输出\\2026-W37")
"""
import argparse
import copy
import hashlib
import os
import random
import re
import shutil
import sys
import zipfile

try:
    from docx import Document
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn, nsdecls
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
    from lxml import etree
except ImportError:
    sys.stderr.write(
        "缺少 python-docx / lxml，请先安装：python -m pip install python-docx\n")
    sys.exit(2)


BASE = r"C:\Users\15258\Desktop\地产周报Agent"
HISTORY_DIR = r"C:\Users\15258\Desktop\实习"
TEMPLATE = os.path.join(BASE, "01_模板", "公司公告周报模板.docx")
WRITABLE_ROOTS = (
    os.path.join(BASE, "04_工作区"),
    os.path.join(BASE, "05_输出"),
)

# (源 md, 输出 docx) —— Word 只有一份完整底稿，没有正式/备选之分
JOBS = [
    ("本周公司公告汇总.md", "本周公司公告周报.docx"),
    ("本周重要资讯.md", "本周重要资讯.docx"),
]

W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"

# "一行空行"的段前间距：260 twips = 13pt，正好是模板日标题自带的分节日距，
# 也约等于 11pt 正文的一行高（11pt × 1.158 = 12.74pt）。
GAP_TWIPS = 260

# w:pPr 里排在 w:contextualSpacing **之后**的元素——插入新元素时要用到
_PPR_AFTER_CONTEXTUAL = (
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
    "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
    "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)

# w:pPr 里排在 w:pageBreakBefore **之后**的元素（ECMA-376 规定的 pPr 子元素顺序）
_PPR_AFTER_PAGEBREAK = (
    "w:framePr", "w:widowControl", "w:numPr", "w:suppressLineNumbers",
    "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
    "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
    "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
    "w:rPr", "w:sectPr", "w:pPrChange",
)

# 日期块标题：先认模板的日标题写法「广发证券房地产每日新闻（2026.09.04）——李怡慧」，
# 再认「标题就是日期本身」的写法（`## 2026.09.04`）。
# 刻意**不**认「含日期但另有主题」的小标题（如 `## 2026年半年度业绩`），
# 否则普通章节会被误判成日期块而多插一个分页。
RE_DATE = re.compile(r"\d{4}\s*[.\-/年]\s*\d{1,2}\s*[.\-/月]\s*\d{1,2}\s*日?")
RE_DATE_ONLY = re.compile(
    r"^[（(]?\s*\d{4}\s*[.\-/年]\s*\d{1,2}\s*[.\-/月]\s*\d{1,2}\s*日?\s*[）)]?"
    r"\s*(?:[—\-–]{1,2}\s*\S+\s*)?$")


def is_date_head(text):
    """该标题是否开启一个新的日期块。"""
    t = (text or "").strip()
    if "每日新闻" in t and RE_DATE.search(t):
        return True
    return bool(RE_DATE_ONLY.match(t))


class WordToolError(Exception):
    """工具层拒绝执行（安全约束或调用错误）。"""


def log(msg=""):
    """GBK 控制台下也不会崩的输出。"""
    try:
        print(msg)
    except UnicodeEncodeError:
        enc = sys.stdout.encoding or "ascii"
        print(msg.encode(enc, "replace").decode(enc, "replace"))


# ----------------------------------------------------------- 路径安全
def _abs(p):
    return os.path.abspath(p)


def _under(path, root):
    p = _abs(path).rstrip("\\/").lower()
    r = _abs(root).rstrip("\\/").lower()
    return p == r or p.startswith(r + os.sep)


def assert_writable_target(out):
    """输出必须落在工作区内，且不得是模板原件或历史实习目录。"""
    if _abs(out).lower() == _abs(TEMPLATE).lower():
        raise WordToolError("拒绝覆盖模板原件：%s" % out)
    if _under(out, HISTORY_DIR):
        raise WordToolError("拒绝写入历史实习目录（只读）：%s" % out)
    if not any(_under(out, r) for r in WRITABLE_ROOTS):
        raise WordToolError(
            "输出路径必须位于 04_工作区 或 05_输出 之内：%s\n  允许的根目录：\n    - %s"
            % (out, "\n    - ".join(WRITABLE_ROOTS)))
    if not out.lower().endswith(".docx"):
        raise WordToolError("输出文件必须是 .docx：%s" % out)


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ----------------------------------------------------------- markdown 解析
RE_TABLE = re.compile(r"^\s*\|.*\|\s*$")
RE_ENTRY = re.compile(r"^(\s*)-\s+(.*)$")
RE_RULE = re.compile(r"^-{3,}$")

# 「模板体」行：不带任何 markdown 标记。本周公司公告正式版.md 通篇是这种格式，
# 它本来就是照 Word 模板的样子写出来的，所以必须单独认出来，否则会被当成
# 普通段落——日期分组会丢掉 16pt 加粗、条目会丢掉 List Paragraph 和编号。
RE_PLAIN_HEAD = re.compile(
    r"^[^【].{0,20}（\s*\d{4}\s*[.\-/年]\s*\d{1,2}\s*[.\-/月]\s*\d{1,2}\s*日?\s*）")
RE_PLAIN_ENTRY = re.compile(r"^【[^】]{1,24}】.+")


def parse_md(text):
    """把 md 拆成块序列：[(kind, text)]，kind ∈
    head1/head2/head3 / news / entry / sub / quote / para / rule / table
    """
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue

        if RE_TABLE.match(line):
            rows = []
            while i < n and RE_TABLE.match(lines[i].rstrip()):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                # 跳过 |---|---| 分隔行
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue

        if s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#"))
            blocks.append(("head%d" % min(lvl, 3), s.lstrip("#").strip()))
            i += 1
            continue

        if s == "【公司新闻】":
            blocks.append(("news", s))
            i += 1
            continue

        if RE_RULE.match(s):
            blocks.append(("rule", ""))
            i += 1
            continue

        if s.startswith(">"):
            # 连续的 > 行在 markdown 里是**同一个**引用块，合并成一段（不插分隔符，
            # 原文件是软换行，断点落在「；」「、」上）
            parts = []
            while i < n and lines[i].strip().startswith(">"):
                parts.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", "".join(parts)))
            continue

        m = RE_ENTRY.match(line)
        if m:
            indent = len(m.group(1).expandtabs(2))
            blocks.append(("entry" if indent < 2 else "sub", m.group(2).strip()))
            i += 1
            continue

        # 缩进且无标记的行 = markdown 软换行的续行，稍后并回上一块
        if line[:1] in (" ", "\t"):
            blocks.append(("cont", s))
            i += 1
            continue

        if RE_PLAIN_HEAD.match(s):
            blocks.append(("head2", s))
            i += 1
            continue

        if RE_PLAIN_ENTRY.match(s):
            blocks.append(("entry", s))
            i += 1
            continue

        blocks.append(("para", s))
        i += 1

    # 把缩进续行并回上一块（软换行不该另起段落，更不该多占一个编号）
    merged = []
    for k, v in blocks:
        if k == "cont" and merged and merged[-1][0] != "table":
            pk, pv = merged[-1]
            merged[-1] = (pk, pv + v)
            continue
        merged.append((k, v))
    return merged


def parse_inline(s):
    """把 `**x**` 拆成 [(文字, 是否加粗)]；非标记字符一个不丢。

    加粗值为 None 表示"沿用模具自带的格式"（正式版条目因此与模板同款）。
    """
    out = []
    for part in re.split(r"(\*\*.+?\*\*)", s):
        if not part:
            continue
        if len(part) > 4 and part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True))
        else:
            out.append((part, None))
    return out


def plain_text(s):
    """去掉行内 markdown 标记后的纯文字（用于校验）。"""
    return "".join(t for t, _ in parse_inline(s))


# ----------------------------------------------------------- 模具与写盘
def _strip_para_ids(el):
    """deepcopy 出来的段落会带重复的 w14:paraId / w14:textId，去掉让 Word 重新分配。"""
    for e in el.iter():
        for a in list(e.attrib):
            if a.startswith(W14) and a[len(W14):] in ("paraId", "textId"):
                del e.attrib[a]


def set_para_text(p_el, segments):
    """把段落内容整体替换为 segments，格式沿用该段第一个 run 的 rPr。"""
    proto = p_el.find(qn("w:r"))
    for r in p_el.findall(qn("w:r")):
        p_el.remove(r)
    if proto is None:          # 空段落：没有可继承的 run，保持为空
        return
    for text, bold in segments:
        r = copy.deepcopy(proto)
        for ch in list(r):
            if ch.tag != qn("w:rPr"):
                r.remove(ch)
        run = Run(r, None)
        run.text = text
        if bold is not None:
            run.bold = bold
        p_el.append(r)


class Renderer(object):
    """用模板自带的段落当模具，往 body 里追加内容。"""

    def __init__(self, doc):
        self.doc = doc
        self.body = doc.element.body
        self.sectPr = self.body.find(qn("w:sectPr"))
        if self.sectPr is None:
            raise WordToolError("模板缺少 w:sectPr，无法安全清空正文")

        paras = doc.paragraphs
        if len(paras) < 5:
            raise WordToolError("模板段落数异常（%d），与预期结构不符" % len(paras))
        self.d_head = paras[0]._p       # 16pt 加粗日标题
        self.d_news = paras[2]._p       # 【公司新闻】
        self.d_entry = paras[4]._p      # List Paragraph 条目（带 numPr）

        # 全文第一段不施加上间距，避免正文顶部凭空多一行
        self._first = True
        self._last = None               # 最近追加的段落（给表格补上方间距用）

        # 结构自检：模具段落必须认得出来，否则宁可报错也不写坏
        if "每日新闻" not in paras[0].text:
            raise WordToolError("模板第 1 段不是日标题：%r" % paras[0].text[:40])
        if paras[2].text.strip() != "【公司新闻】":
            raise WordToolError("模板第 3 段不是【公司新闻】：%r" % paras[2].text[:40])
        if paras[4].style.name != "List Paragraph":
            raise WordToolError("模板第 5 段不是 List Paragraph：%r" % paras[4].style.name)

        self._nb = None
        self._next_abs = self._next_num = None
        self.date_blocks = 0        # 本次渲染遇到的日期块数（render 结束时回填）

    # -- 编号：每个日期分组一个新列表实例，编号从 1 重启 ---------------
    def _numbering(self):
        if self._nb is not None:
            return self._nb
        try:
            part = self.doc.part.numbering_part
        except (KeyError, NotImplementedError):
            part = None
        if part is None or part.element is None:
            raise WordToolError("模板没有 numbering.xml，无法还原编号格式")
        nb = part.element
        self._nb = nb
        self._next_abs = 1 + max(
            int(x) for x in re.findall(r'<w:abstractNum w:abstractNumId="(\d+)"',
                                       nb.xml.encode("utf-8").decode("utf-8")))
        self._next_num = 1 + max(
            int(x) for x in re.findall(r'<w:num w:numId="(\d+)"',
                                       nb.xml.encode("utf-8").decode("utf-8")))
        return nb

    def new_num_id(self):
        """注册一份新的 abstractNum + numId，返回新 numId（编号重启）。"""
        nb = self._numbering()
        src = nb.find(qn("w:abstractNum"))
        if src is None:
            raise WordToolError("numbering.xml 里没有 abstractNum 可克隆")

        aid = self._next_abs
        self._next_abs += 1

        new_abs = copy.deepcopy(src)
        new_abs.set(qn("w:abstractNumId"), str(aid))
        # nsid / tmpl 是"列表模板"指纹，必须换掉，否则 Word 会把两份定义
        # 认成同一个列表而继续编号
        for tag in ("w:nsid", "w:tmpl"):
            el = new_abs.find(qn(tag))
            if el is not None:
                el.set(qn("w:val"), "%08X" % random.randint(0, 0xFFFFFFFF))

        first_num = nb.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(new_abs)
        else:
            nb.append(new_abs)

        nid = self._next_num
        self._next_num += 1
        new_num = OxmlElement("w:num")
        new_num.set(qn("w:numId"), str(nid))
        aid_el = OxmlElement("w:abstractNumId")
        aid_el.set(qn("w:val"), str(aid))
        new_num.append(aid_el)
        nb.append(new_num)
        return nid

    # -- 间距：所有层级统一"一行空行"，用段前间距实现，不插空段落 --------
    def _apply_gap(self, el):
        pPr = el.get_or_add_pPr()

        # List Paragraph 样式带 <w:contextualSpacing/>，会吞掉相邻同款段落之间的
        # 间距（公告会全部粘在一起）。必须在段落级显式覆写为 0。
        cs = pPr.find(qn("w:contextualSpacing"))
        if cs is None:
            cs = OxmlElement("w:contextualSpacing")
            pPr.insert_element_before(cs, *_PPR_AFTER_CONTEXTUAL)
        cs.set(qn("w:val"), "0")

        # get_or_add_spacing 会插到 w:ind / w:contextualSpacing 之前，顺序合法；
        # 已存在的 spacing（日标题的 before=260 after=260 line=416）只改 before/after，
        # line / lineRule 原样保留 → 模板行距不变。
        sp = pPr.get_or_add_spacing()
        sp.set(qn("w:before"), "0" if self._first else str(GAP_TWIPS))
        sp.set(qn("w:after"), "0")
        self._first = False
        self._last = el

    def gap_after_last(self, twips=GAP_TWIPS):
        """给上一段补**段后**间距——表格不参与段前间距，只能从上方那段补一行。"""
        if self._last is None:
            return
        sp = self._last.get_or_add_pPr().get_or_add_spacing()
        sp.set(qn("w:after"), str(twips))

    # -- 分页：日期块各自另起一页 -------------------------------------
    def set_page_break_before(self, el):
        """让该段**从新的一页顶部开始**（真正的 Word 分页，不是空行凑的）。

        用 `<w:pageBreakBefore/>`（段落属性「段前分页」）而不是手工分页符
        `<w:br w:type="page"/>`，原因有二：
          1. Word 在"本段已经位于页首"时会自动忽略它，**不会凭空多出一张空白页**；
             手工分页符落在页首时是会真的翻出一页空白的。
          2. 它只是段落属性，不新增 run、不新增段落、不产生任何可见字符，
             段落的文字 / 字体 / 编号 / 间距全部不受影响。
        两者都是 Word 原生分页，都不涉及空行——满足"必须使用真正的分页符"。
        """
        pPr = el.get_or_add_pPr()
        pbb = pPr.find(qn("w:pageBreakBefore"))
        if pbb is None:
            pbb = OxmlElement("w:pageBreakBefore")
            pPr.insert_element_before(pbb, *_PPR_AFTER_PAGEBREAK)
        # 与 Word 自己的写法一致：勾选时写裸元素，取消勾选才写 w:val="0"
        if pbb.get(qn("w:val")) is not None:
            del pbb.attrib[qn("w:val")]
        return pbb

    # -- 追加段落 -----------------------------------------------------
    def _append(self, donor, segments):
        el = copy.deepcopy(donor)
        _strip_para_ids(el)
        set_para_text(el, segments)
        self._apply_gap(el)
        self.sectPr.addprevious(el)
        return el

    def head(self, text):
        return self._append(self.d_head, parse_inline(text))

    def date_head(self, text, new_page=True):
        """日期块标题：需要时让它另起一页（第一个日期块 new_page=False）。"""
        el = self.head(text)
        if new_page:
            self.set_page_break_before(el)
        return el

    def news(self, text):
        return self._append(self.d_news, parse_inline(text))

    def entry(self, text, num_id):
        el = self._append(self.d_entry, parse_inline(text))
        pPr = el.get_or_add_pPr()
        numPr = pPr.get_or_add_numPr()
        numPr.get_or_add_numId().val = num_id
        numPr.get_or_add_ilvl().val = 0
        return el

    def body_para(self, text, indent=None):
        el = self._append(self.d_news, parse_inline(text))
        if indent:
            pPr = el.get_or_add_pPr()
            pPr.get_or_add_ind().set(qn("w:left"), str(indent))
        return el

    # -- 表格 ---------------------------------------------------------
    def table(self, rows):
        if not rows:
            return None
        ncols = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=ncols)

        tblPr = t._tbl.tblPr
        borders = parse_xml(
            '<w:tblBorders %s>'
            '<w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            '<w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            '<w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            '</w:tblBorders>' % nsdecls("w"))
        tblPr.insert_element_before(
            borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
            "w:tblCaption", "w:tblDescription", "w:tblPrChange")

        for ri, row in enumerate(rows):
            for ci in range(ncols):
                cell = t.cell(ri, ci)
                txt = row[ci] if ci < len(row) else ""
                p = cell.paragraphs[0]
                run = p.add_run(txt)
                run.font.name = "Calibri"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                if ri == 0:
                    run.bold = True
        self.sectPr.addprevious(t._tbl)
        self._last = None          # 表格后面那段自己的段前间距负责留白
        return t


# ----------------------------------------------------------- 渲染主流程
def render(doc, blocks):
    r = Renderer(doc)

    # 清空正文（保留 sectPr）——模板里的旧公告内容在这一步被清掉
    for child in list(r.body):
        if child is not r.sectPr:
            r.body.remove(child)

    # 空行一律由 Renderer._apply_gap 的段前间距产生，这里不再插入任何空段落——
    # 所有层级（日标题 / 【公司新闻】 / 条目之间 / 日块之间）因此都是同一行距。
    i = 0
    n = len(blocks)
    date_blocks = 0          # 已开始的日期块数
    while i < n:
        kind, val = blocks[i]

        if kind in ("head1", "head2", "head3", "news"):
            if kind != "news" and is_date_head(val):
                date_blocks += 1
                # 第 1 个日期块不加分页（第一页本来就在页首）；
                # 之后每进入一个新日期，都让它从新的一页顶部开始，
                # 因此新日期标题永远不会落在上一页末尾。
                r.date_head(val, new_page=(date_blocks > 1))
            elif kind == "news":
                r.news(val)
            else:
                r.head(val)
            i += 1
            continue

        if kind == "entry":
            group = []
            while i < n and blocks[i][0] in ("entry", "sub"):
                group.append(blocks[i])
                i += 1
            num_id = r.new_num_id()          # 本组编号从 1 重启
            for k, t in group:
                if k == "entry":
                    r.entry(t, num_id)
                else:
                    r.body_para(t, indent=1080)   # 子项：缩进、不占编号
            continue

        if kind == "table":
            r.gap_after_last()               # 表格上方补一行（见 gap_after_last）
            r.table(val)
            i += 1
            continue

        if kind == "rule":
            i += 1
            continue

        if kind == "quote":
            r.body_para(val, indent=720)
            i += 1
            continue

        r.body_para(val)
        i += 1

    r.date_blocks = date_blocks
    return r


def generate(out_dir, force=False, template=TEMPLATE, jobs=None):
    """生成 docx（公告周报 + 重要资讯）。返回 [(输出路径, 段落数)]。"""
    out_dir = _abs(out_dir)
    if not os.path.isfile(template):
        raise WordToolError("模板不存在：%s" % template)
    tpl_sha = sha256(template)

    jobs = jobs or JOBS
    missing = [src for src, _ in jobs
               if not os.path.isfile(os.path.join(out_dir, src))]
    if missing:
        raise WordToolError("源 md 不存在：\n    - %s" % "\n    - ".join(
            os.path.join(out_dir, m) for m in missing))

    results = []
    for src_name, out_name in jobs:
        src = os.path.join(out_dir, src_name)
        out = os.path.join(out_dir, out_name)
        assert_writable_target(out)
        if os.path.exists(out) and not force:
            raise WordToolError(
                "目标文件已存在，拒绝覆盖：%s\n  如需覆盖请显式加 --force" % out)

        # 1) 复制模板 → 副本；模板原件此后不再被打开
        shutil.copyfile(template, out)
        # 2) 打开副本修改
        doc = Document(out)
        blocks = parse_md(open(src, encoding="utf-8").read())
        r = render(doc, blocks)
        doc.save(out)

        # 3) 模板必须原封不动
        if sha256(template) != tpl_sha:
            raise WordToolError("模板被改动了！%s" % template)

        results.append((out, len(doc.paragraphs), src))
        counts = {}
        for k, _ in blocks:
            counts[k] = counts.get(k, 0) + 1
        log("  [OK] %-24s <- %s" % (out_name, src_name))
        log("       分块：%s" % "  ".join(
            "%s=%d" % (k, v) for k, v in sorted(counts.items())))
        if r.date_blocks:
            log("       分页：%d 个日期块 → %d 个分页符（首个日期块不加）"
                % (r.date_blocks, r.date_blocks - 1))
        else:
            log("       分页：无日期块，未插入分页符")
    return results


# ----------------------------------------------------------- 校验
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _canon(el):
    """元素的规范化字节：忽略属性顺序、命名空间前缀写法与空白。

    python-docx 存盘时会重排整份 XML 的序列化，所以**不能**拿字节比对——
    字节不同不代表内容被改了。规范化之后比对才是"语义是否相同"。
    """
    try:
        return etree.tostring(el, method="c14n2", with_comments=False)
    except (ValueError, TypeError):
        return etree.tostring(el, method="c14n", with_comments=False)


def _template_fidelity(template, out, problems):
    """核对模板原有的**字体 / 字号 / 编号 / 页边距 / 页眉页脚 / 段落样式**未被改动。

    直接比对整包字节是不行的（见 `_canon`），因此改为逐项核对真正决定版式的四处：
      `styles.xml`      → 段落样式 / 字体 / 字号
      `header*` `footer*` → 页眉页脚（这两类部件我们从不触碰，要求逐字节相同）
      `w:sectPr`        → 纸张与页边距
      `abstractNum`     → 编号格式（模板原有的定义必须一个不少、内容不变）
    """
    zt, zo = zipfile.ZipFile(template), zipfile.ZipFile(out)
    tn, on = set(zt.namelist()), set(zo.namelist())

    for name in sorted(tn):
        base = name.rsplit("/", 1)[-1]
        hf = base.startswith("header") or base.startswith("footer")
        if name != "word/styles.xml" and not hf:
            continue
        if name not in on:
            problems.append("模板部件丢失：%s" % name)
            continue
        a, b = zt.read(name), zo.read(name)
        if hf:
            if a != b:
                problems.append("页眉/页脚被改动：%s" % name)
        elif _canon(etree.fromstring(a)) != _canon(etree.fromstring(b)):
            problems.append("样式表被改动（字体/字号/段落样式可能已变）：%s" % name)

    def _sectpr(z):
        return etree.fromstring(z.read("word/document.xml")).find(
            ".//" + W_NS + "sectPr")

    sa, sb = _sectpr(zt), _sectpr(zo)
    if sa is None or sb is None:
        problems.append("document.xml 里找不到 w:sectPr，无法核对页面设置")
    elif _canon(sa) != _canon(sb):
        problems.append("页面设置（纸张 / 页边距）被改动")

    def _abstract_nums(z):
        return {e.get(W_NS + "abstractNumId"): _canon(e)
                for e in etree.fromstring(z.read("word/numbering.xml")).iter(
                    W_NS + "abstractNum")}

    na, nb = _abstract_nums(zt), _abstract_nums(zo)
    for k, v in na.items():
        if k not in nb:
            problems.append("模板的编号定义 abstractNumId=%s 丢失" % k)
        elif nb[k] != v:
            problems.append("模板的编号定义 abstractNumId=%s 被改动" % k)


def verify(out_dir, jobs=None, template=TEMPLATE):
    """读回 docx：zip 完整性 / 内容与 md 逐条一致 / 排版硬规则 / 模板保真。"""
    jobs = jobs or JOBS
    ok = True
    for src_name, out_name in jobs:
        src = os.path.join(out_dir, src_name)
        out = os.path.join(out_dir, out_name)
        problems = []

        # a) zip 完整
        if not zipfile.is_zipfile(out):
            problems.append("不是合法的 zip/docx")
        else:
            try:
                z = zipfile.ZipFile(out)
                bad = z.testzip()
                if bad:
                    problems.append("zip 损坏：%s" % bad)
                for name in ("word/document.xml", "word/styles.xml",
                             "word/numbering.xml"):
                    if name not in z.namelist():
                        problems.append("缺少 %s" % name)
            except Exception as exc:                       # noqa: BLE001
                problems.append("zip 读取失败：%s" % exc)

        # b) 文件能被 Word 解析器打开、且内容与 md 逐条对齐
        try:
            d = Document(out)
        except Exception as exc:                           # noqa: BLE001
            problems.append("python-docx 打不开：%s" % exc)
            d = None

        if d is not None:
            # 按 body 真实子元素顺序取内容（表格必须落在它该在的位置，
            # 不能把表格单元格一股脑追加到末尾——那样测不出顺序错乱）
            got = []
            for ch in d.element.body:
                if ch.tag == qn("w:p"):
                    t = Paragraph(ch, d).text
                    if t.strip():
                        got.append(t)
                elif ch.tag == qn("w:tbl"):
                    for row in Table(ch, d).rows:
                        got.append(" | ".join(c.text for c in row.cells))
            if len(d.tables) == 0 and any(k == "table" for k, _ in parse_md(
                    open(src, encoding="utf-8").read())):
                problems.append("md 有表格，docx 里没有")

            src_blocks = parse_md(open(src, encoding="utf-8").read())
            want = []
            for k, v in src_blocks:
                if k == "table":
                    want += [" | ".join(plain_text(c) for c in row)
                             for row in v]
                elif plain_text(v).strip():
                    want.append(plain_text(v))

            if got != want:
                problems.append(
                    "内容与 md 不是逐条逐序一致（docx %d 条 / md %d 条）"
                    % (len(got), len(want)))
                for i, (a, b) in enumerate(zip(got, want)):
                    if a != b:
                        problems.append("  首个差异 @%d docx=%r md=%r"
                                        % (i, a[:40], b[:40]))
                        break

            # c) 公告条目必须真的是 List Paragraph 且带编号
            #    （防"模板体"行被当成普通段落——那会丢掉编号和加粗）
            entry_texts = set(plain_text(v) for k, v in src_blocks if k == "entry")
            # 只有日期/章节分组标题才是 16pt 加粗；【公司新闻】在模板里是普通段落
            flat = [plain_text(v) for k, v in src_blocks
                    if k in ("head1", "head2", "head3")]
            for p in d.paragraphs:
                if p.text in entry_texts:
                    pPr = p._element.pPr
                    has_num = pPr is not None and pPr.find(qn("w:numPr")) is not None
                    if p.style.name != "List Paragraph" or not has_num:
                        problems.append(
                            "条目缺 List Paragraph/编号：%r（样式=%s，编号=%s）"
                            % (p.text[:26], p.style.name, has_num))
                    # List Paragraph 自带的 contextualSpacing 若没被覆写，
                    # 相邻条目之间不会留白 → 公告会粘成一片
                    cs = pPr.find(qn("w:contextualSpacing")) if pPr is not None else None
                    if cs is None or cs.get(qn("w:val")) not in ("0", "false"):
                        problems.append(
                            "条目未覆写 contextualSpacing（会与下一条粘住）：%r"
                            % p.text[:26])
                if p.text in flat and p.text.strip():
                    r0 = p.runs[0] if p.runs else None
                    if not (r0 and r0.font.size and r0.font.bold):
                        problems.append(
                            "分组标题不是 16pt 加粗：%r" % p.text[:26])
            # d) 排版硬规则：全文不插空段落，所有层级统一"一行空行"
            body_paras = [Paragraph(ch, d)
                          for ch in d.element.body if ch.tag == qn("w:p")]
            blanks = [i for i, p in enumerate(body_paras) if not p.text.strip()]
            if blanks:
                problems.append(
                    "出现 %d 个空段落（空行必须由段前间距产生，不能手插空行）：位置 %s"
                    % (len(blanks), blanks[:6]))
            for i, p in enumerate(body_paras):
                pPr = p._element.pPr
                sp = pPr.find(qn("w:spacing")) if pPr is not None else None
                if sp is None:
                    problems.append("第 %d 段没有显式 spacing：%r" % (i, p.text[:24]))
                    continue
                before = int(sp.get(qn("w:before")) or 0)
                after = int(sp.get(qn("w:after")) or 0)
                want = 0 if i == 0 else GAP_TWIPS
                if before != want or after != 0:
                    problems.append(
                        "第 %d 段间距不是统一一行空行：before=%d after=%d（应为 %d/0）：%r"
                        % (i, before, after, want, p.text[:24]))

            # e) 分页硬规则：每个日期块另起一页（首个日期块除外，第一页不加）
            date_heads = [plain_text(v) for k, v in src_blocks
                          if k in ("head1", "head2", "head3") and is_date_head(v)]
            want_pb = max(0, len(date_heads) - 1)
            got_pb = []
            for j, p in enumerate(body_paras):
                pPr = p._element.pPr
                pbb = pPr.find(qn("w:pageBreakBefore")) if pPr is not None else None
                if pbb is not None and pbb.get(qn("w:val")) not in ("0", "false", "off"):
                    got_pb.append(j)
            if len(got_pb) != want_pb:
                problems.append(
                    "分页符数量不对：docx %d 个 / md %d 个日期块 → 应为 %d 个"
                    % (len(got_pb), len(date_heads), want_pb))
            for j in got_pb:
                if not is_date_head(body_paras[j].text):
                    problems.append(
                        "分页符落在了非日期块标题上（第 %d 段：%r）"
                        % (j, body_paras[j].text[:30]))
            if date_heads:
                first_head = date_heads[0]
                for j in got_pb:
                    if body_paras[j].text.strip() == first_head.strip():
                        problems.append(
                            "首个日期块不应加段前分页（第一页本来就在页首）：%r"
                            % first_head[:30])
            # 分页必须落在日期块标题上、且不得用空行凑
            if got_pb and blanks:
                problems.append("同时存在分页符与空段落，分页可能由空行凑出")

            if not entry_texts:
                problems.append("md 里一条公告条目都没识别出来（分块规则可能失效）")

            # f) 模板保真：字体/字号/编号/页边距/页眉页脚/段落样式必须原样保留
            if template and os.path.isfile(template):
                try:
                    _template_fidelity(template, out, problems)
                except Exception as exc:                   # noqa: BLE001
                    problems.append("与模板比对失败：%s" % exc)

        log("  [%s] %-24s 段落 %s，表格 %s"
            % ("OK" if not problems else "FAIL", out_name,
               len(d.paragraphs) if d else "-",
               len(d.tables) if d else "-"))
        for p in problems:
            ok = False
            log("         - %s" % p)
    return ok


def pick_jobs(only):
    """把 --only 解析成 JOBS 的子集。

    支持传输出名 / 源 md 名 / 二者中的关键字（如「汇总」「重要资讯」）——
    用户说"只重新生成某一份"时不必记全名。
    """
    if not only:
        return list(JOBS)
    q = only.strip()
    hits = [j for j in JOBS if q in (j[0], j[1]) or q in j[0] or q in j[1]]
    if len(hits) != 1:
        raise WordToolError(
            "--only %r 匹配到 %d 项，请写得更明确。可选：\n    - %s"
            % (only, len(hits), "\n    - ".join(o for _, o in JOBS)))
    return hits


def current_week_dir():
    root = os.path.join(BASE, "05_输出")
    subs = [d for d in os.listdir(root) if os.path.isdir(os.path.join(root, d))]
    if not subs:
        raise WordToolError("05_输出 下没有周次目录")
    return os.path.join(root, sorted(subs)[-1])


def main():
    ap = argparse.ArgumentParser(
        description="把周报 md 交付物按模板渲染成 docx（模板原件只读）")
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--week", help="周次，例如 2026-W37（等价于 --outdir 05_输出\\2026-W37）")
    g.add_argument("--outdir", help="直接指定输出目录")
    ap.add_argument("--force", action="store_true", help="允许覆盖已存在的 docx")
    ap.add_argument("--verify-only", action="store_true", help="只校验，不生成")
    ap.add_argument("--only", help="只处理其中一项，例如 --only 本周公司公告周报.docx "
                                   "或 --only 汇总；不填则两项都处理")
    ap.add_argument("--template", default=TEMPLATE)
    args = ap.parse_args()

    if args.outdir:
        out_dir = _abs(args.outdir)
    elif args.week:
        out_dir = os.path.join(BASE, "05_输出", args.week)
    else:
        out_dir = current_week_dir()

    log("模板  : %s" % _abs(args.template))
    log("目录  : %s" % out_dir)
    log("")

    try:
        jobs = pick_jobs(args.only)
        if args.only:
            log("范围  : 只处理 %s" % ", ".join(o for _, o in jobs))
            log("")
        if not args.verify_only:
            log("生成：")
            generate(out_dir, force=args.force, template=args.template, jobs=jobs)
            log("")
        log("校验：")
        ok = verify(out_dir, jobs=jobs, template=args.template)
    except WordToolError as exc:
        log("[拒绝] %s" % exc)
        return 3

    log("")
    log("模板 sha256 未变：%s" % sha256(args.template))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
