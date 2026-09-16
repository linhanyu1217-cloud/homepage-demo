# -*- coding: utf-8 -*-
"""Excel 结构化更新的**基础**封装（openpyxl 3.1.5）。

当前阶段只提供最小能力，**不绑定任何具体业务字段**：
    打开 xlsx / 列出 sheet / 读单元格 / 读区域 / 写指定单元格 / 另存为新文件

安全约定（硬规则，代码强制，不可用参数绕过）：
    1. **默认不覆盖原文件** —— 写入一律走「另存为新文件」，目标路径不能等于源路径；
    2. **默认只允许写到** 04_工作区 或 05_输出（工作区内部）；
    3. **禁止写入/修改** 历史实习目录 `C:\\Users\\15258\\Desktop\\实习` 下的任何 Excel；
    4. 目标文件已存在时**拒绝覆盖**，除非显式 `--force`；
    5. 不注册任何业务字段、不做自动列映射 —— 由调用方（Agent 或人工）显式指定。

命令行用法：
    python 脚本\\excel工具.py info   <源文件>
    python 脚本\\excel工具.py sheets <源文件>
    python 脚本\\excel工具.py read   <源文件> <sheet> <A1 或 A1:C10>
    python 脚本\\excel工具.py copy   <源文件> --out 04_工作区\\副本.xlsx
    python 脚本\\excel工具.py write  <源文件> <sheet> <单元格> <值> --out 04_工作区\\新.xlsx

作为模块导入：
    from excel工具 import open_book, list_sheets, read_cell, write_cell, save_as
"""
import argparse
import datetime as dt
import os
import sys

try:
    from openpyxl import load_workbook
    from openpyxl.utils import range_boundaries
except ImportError:
    sys.stderr.write("缺少 openpyxl，请先安装：python -m pip install openpyxl\n")
    sys.exit(2)

BASE = r"C:\Users\15258\Desktop\地产周报Agent"
HISTORY_DIR = r"C:\Users\15258\Desktop\实习"
WRITABLE_ROOTS = (
    os.path.join(BASE, "04_工作区"),
    os.path.join(BASE, "05_输出"),
)


class ExcelToolError(Exception):
    """工具层拒绝执行（安全约束或调用错误）。"""


def log(msg=""):
    try:
        print(msg)
    except UnicodeEncodeError:
        enc = sys.stdout.encoding or "ascii"
        print(msg.encode(enc, "replace").decode(enc, "replace"))


# ----------------------------------------------------------- 路径安全
def _abs(p):
    return os.path.abspath(p)


def _under(path, root):
    """path 是否在 root 之内（含 root 本身）。"""
    p = _abs(path).rstrip("\\/").lower()
    r = _abs(root).rstrip("\\/").lower()
    return p == r or p.startswith(r + os.sep)


def assert_readable(src):
    """源文件可读即可（历史实习目录允许**读**，只是禁止写）。"""
    if not os.path.isfile(src):
        raise ExcelToolError("源文件不存在：%s" % src)
    if not src.lower().endswith((".xlsx", ".xlsm")):
        raise ExcelToolError("只支持 .xlsx / .xlsm：%s" % src)


def assert_writable_target(out, src=None):
    """写入目标必须落在工作区内的可写目录，且不等于源文件。"""
    if not out:
        raise ExcelToolError("必须显式指定输出路径 --out（默认不覆盖原文件）")

    if _under(out, HISTORY_DIR):
        raise ExcelToolError(
            "拒绝写入历史实习目录（只读）：%s" % out)

    if not any(_under(out, r) for r in WRITABLE_ROOTS):
        raise ExcelToolError(
            "输出路径必须位于 04_工作区 或 05_输出 之内：%s\n"
            "  允许的根目录：\n    - %s"
            % (out, "\n    - ".join(WRITABLE_ROOTS)))

    if src and _abs(out).lower() == _abs(src).lower():
        raise ExcelToolError("拒绝覆盖源文件：%s" % out)

    if not out.lower().endswith((".xlsx", ".xlsm")):
        raise ExcelToolError("输出文件必须是 .xlsx / .xlsm：%s" % out)


# ----------------------------------------------------------- 基础操作
def open_book(path, data_only=False):
    """打开工作簿。

    data_only=False（默认）保留公式；True 读取公式的**缓存值**（无缓存则为 None）。
    """
    assert_readable(path)
    try:
        return load_workbook(filename=path, data_only=data_only)
    except Exception as exc:                       # noqa: BLE001
        raise ExcelToolError("打开失败：%s（%s）" % (path, exc))


def list_sheets(wb):
    """返回 sheet 名称列表。"""
    return list(wb.sheetnames)


def get_sheet(wb, name=None):
    """取 sheet；name 为空则取活动表。找不到 → 报错（不静默 fallback）。"""
    if name is None or name == "":
        return wb.active
    if name not in wb.sheetnames:
        raise ExcelToolError(
            "工作表不存在：%r\n  可用的有：%s" % (name, "、".join(wb.sheetnames)))
    return wb[name]


def read_cell(wb, sheet, ref, data_only=False):
    """读单个单元格。

    若已用 data_only=True 打开，传入的 wb 已带缓存值，这里不再重复加载。
    """
    ws = get_sheet(wb, sheet)
    try:
        return ws[ref].value
    except Exception as exc:                       # noqa: BLE001
        raise ExcelToolError("读取单元格失败 %s!%s（%s）" % (sheet, ref, exc))


def read_range(wb, sheet, ref):
    """读区域，返回二维列表（行优先）。单格也返回 [[值]]。"""
    ws = get_sheet(wb, sheet)
    try:
        min_c, min_r, max_c, max_r = range_boundaries(ref)
    except Exception as exc:                       # noqa: BLE001
        raise ExcelToolError("区域格式不合法：%r（%s）" % (ref, exc))
    if min_r == max_r and min_c == max_c:
        return [[ws.cell(row=min_r, column=min_c).value]]
    return [[ws.cell(row=r, column=c).value
             for c in range(min_c, max_c + 1)]
            for r in range(min_r, max_r + 1)]


def write_cell(wb, sheet, ref, value):
    """写单个单元格（只改内存中的工作簿，落盘需再调用 save_as）。"""
    ws = get_sheet(wb, sheet)
    try:
        ws[ref] = value
    except Exception as exc:                       # noqa: BLE001
        raise ExcelToolError("写入单元格失败 %s!%s（%s）" % (sheet, ref, exc))
    return ws


def save_as(wb, out, src=None, force=False):
    """另存为新文件。默认拒绝覆盖已存在的文件（除非 force=True）。"""
    assert_writable_target(out, src=src)
    if os.path.exists(out) and not force:
        raise ExcelToolError(
            "目标文件已存在，拒绝覆盖：%s\n  如需覆盖请显式加 --force" % out)
    os.makedirs(os.path.dirname(_abs(out)), exist_ok=True)
    try:
        wb.save(out)
    except Exception as exc:                       # noqa: BLE001
        raise ExcelToolError("保存失败：%s（%s）" % (out, exc))
    return out


def copy_to(src, out, force=False):
    """把源工作簿另存为副本。源文件只读，绝不被改动。"""
    wb = open_book(src)
    return save_as(wb, out, src=src, force=force)


# ----------------------------------------------------------- 命令行
def _coerce(text, kind):
    if kind == "str" or text is None:
        return text
    if kind == "int":
        return int(text)
    if kind == "float":
        return float(text)
    # auto：能转成数字就转，否则原样字符串
    try:
        return int(text)
    except (TypeError, ValueError):
        pass
    try:
        return float(text)
    except (TypeError, ValueError):
        return text


def _fmt(v):
    if v is None:
        return "(空)"
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, dt.date):
        return v.isoformat()
    return str(v)


def main():
    ap = argparse.ArgumentParser(
        description="Excel 基础工具（openpyxl）—— 默认不覆盖原文件")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("info", help="显示工作簿概况")
    p.add_argument("src")

    p = sub.add_parser("sheets", help="列出所有 sheet")
    p.add_argument("src")

    p = sub.add_parser("read", help="读取单元格或区域")
    p.add_argument("src")
    p.add_argument("sheet", nargs="?", default=None)
    p.add_argument("ref")
    p.add_argument("--data-only", action="store_true",
                   help="读取公式的缓存值而非公式本身")

    p = sub.add_parser("copy", help="另存为副本（源文件不动）")
    p.add_argument("src")
    p.add_argument("--out", required=True)
    p.add_argument("--force", action="store_true")

    p = sub.add_parser("write", help="写指定单元格后另存为新文件")
    p.add_argument("src")
    p.add_argument("sheet", nargs="?", default=None)
    p.add_argument("ref")
    p.add_argument("value")
    p.add_argument("--out", required=True)
    p.add_argument("--type", choices=["auto", "str", "int", "float"],
                   default="auto")
    p.add_argument("--force", action="store_true")

    args = ap.parse_args()

    try:
        if args.cmd == "info":
            wb = open_book(args.src)
            log("文件      : %s" % _abs(args.src))
            log("大小      : %.1f KB" % (os.path.getsize(args.src) / 1024.0))
            log("sheet 数  : %d" % len(wb.sheetnames))
            for name in wb.sheetnames:
                ws = wb[name]
                log("  - %-24s %d 行 x %d 列"
                    % (name, ws.max_row or 0, ws.max_column or 0))
            return 0

        if args.cmd == "sheets":
            wb = open_book(args.src)
            for name in list_sheets(wb):
                log(name)
            return 0

        if args.cmd == "read":
            wb = open_book(args.src, data_only=args.data_only)
            ws = get_sheet(wb, args.sheet)
            rows = read_range(wb, args.sheet, args.ref)
            log("# %s!%s" % (ws.title, args.ref))
            for r in rows:
                log(" | ".join(_fmt(v) for v in r))
            return 0

        if args.cmd == "copy":
            out = copy_to(args.src, args.out, force=args.force)
            log("已另存    : %s" % _abs(out))
            log("源文件未动: %s" % _abs(args.src))
            _warn_lossy()
            return 0

        if args.cmd == "write":
            wb = open_book(args.src)
            ws = get_sheet(wb, args.sheet)
            old = ws[args.ref].value
            val = _coerce(args.value, args.type)
            write_cell(wb, ws.title, args.ref, val)
            out = save_as(wb, args.out, src=args.src, force=args.force)
            log("已写入    : %s!%s  %s -> %s"
                % (ws.title, args.ref, _fmt(old), _fmt(val)))
            log("已另存    : %s" % _abs(out))
            log("源文件未动: %s" % _abs(args.src))
            _warn_lossy()
            return 0

    except ExcelToolError as exc:
        log("[拒绝] %s" % exc)
        return 3
    return 0


def _warn_lossy():
    log("")
    log("[注意] openpyxl 重新保存会丢失：图表、数据透视表、宏(VBA)、"
        "部分条件格式与图片。")
    log("       若源文件含上述内容，请勿用本工具处理，改由人工另存。")


if __name__ == "__main__":
    sys.exit(main())
