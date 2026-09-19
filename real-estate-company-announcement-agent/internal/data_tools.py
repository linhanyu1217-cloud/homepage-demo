# -*- coding: utf-8 -*-
r"""数据工具集（由 读取文档 / 合并公告汇总 / 盘点公司公告 / 构建公司池 /
提取港股公司 / 提取本周文本 / excel工具 / 检查依赖 八个脚本整合而来）。

本文件是**内部实现**，使用者不应直接调用；统一走 internal\agent_core.py 的
公开函数：

    read_documents()                    读取 docx/pptx/xlsx/pdf 纯文本
    merge_announcements()               正式版 + 备选 合并为 Word 中间底稿
    build_company_pool()                盘点公告目录 → 行业公司池.csv
    extract_hk_companies()              从港股公告提取公司名与代码
    extract_week_text()                 按周次批量抽取 PDF/docx 正文
    check_dependencies()                检查 Python 依赖（只报告，绝不自动安装）
    excel_*()                           Excel 读写（openpyxl 封装）

路径约定：全部相对项目根，靠 __file__ 推导。
HISTORY_DIR 由 agent_core 按 00_配置\个人配置.json 注入；发行版为 None，
安全护栏自动退化为「不检查该项」。
"""
import argparse
import csv
import datetime as dt
import importlib
import json
import os
import re
import sys
import zipfile
from collections import defaultdict
from xml.etree import ElementTree as ET

try:
    from openpyxl import load_workbook
    from openpyxl.utils import range_boundaries
except ImportError:
    sys.stderr.write("缺少 openpyxl，请先安装：python -m pip install openpyxl\n")
    sys.exit(2)

# --------------------------------------------------------------- 路径（相对化）
_HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(_HERE)                       # 项目根目录
BASE = ROOT                                         # 兼容旧脚本里的 BASE 命名

OUT_ROOT = os.path.join(ROOT, "05_输出")
WORK_ROOT = os.path.join(ROOT, "04_工作区")
PERSONAL_DIR = os.path.join(ROOT, "03_个人资料")
INPUT_ROOT = os.path.join(ROOT, "02_待处理")

# 盘点 / 港股公司提取的中间产物目录（A股公司统计.csv、港股公司提取.csv 等）。
# 原 构建公司池.py 里写作 PAND，这里沿用该名，供 load_ashare() / load_hk() 使用。
PAND = os.path.join(WORK_ROOT, "盘点")

WRITABLE_ROOTS = (WORK_ROOT, OUT_ROOT)
HISTORY_DIR = None          # 个人版由 agent_core 注入；发行版保持 None


def log(msg=""):
    """GBK 控制台下也不会崩的输出。"""
    try:
        print(msg)
    except UnicodeEncodeError:
        enc = sys.stdout.encoding or "ascii"
        print(msg.encode(enc, "replace").decode(enc, "replace"))


def _abs(p):
    return os.path.abspath(p)


def _under(path, root):
    # root 为空（发行版没有历史实习目录）时视为「不可能落在其下」
    if not root:
        return False
    p = _abs(path).rstrip("\\/").lower()
    r = _abs(root).rstrip("\\/").lower()
    return p == r or p.startswith(r + os.sep)


def history_dir():
    """当前生效的历史实习资料目录（可能为 None）。"""
    return HISTORY_DIR


# ==================================================================
#  1. 文档读取（原 读取文档.py）
# ==================================================================
NS = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'ss': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
}


def _text_of(elem, tag):
    return ''.join(t.text or '' for t in elem.iter(tag))


def read_docx(path, max_chars):
    out = []
    with zipfile.ZipFile(path) as z:
        with z.open('word/document.xml') as f:
            root = ET.parse(f).getroot()
        for p in root.iter('{%s}p' % NS['w']):
            line = _text_of(p, '{%s}t' % NS['w']).strip()
            if line:
                out.append(line)
            if sum(len(x) for x in out) > max_chars:
                out.append('...[截断]')
                break
    return '\n'.join(out)


def read_pptx(path, max_chars):
    out = []
    with zipfile.ZipFile(path) as z:
        slides = sorted(
            (n for n in z.namelist()
             if re.match(r'ppt/slides/slide\d+\.xml$', n)),
            key=lambda n: int(re.search(r'(\d+)', n.split('/')[-1]).group(1)),
        )
        total = 0
        for name in slides:
            with z.open(name) as f:
                root = ET.parse(f).getroot()
            texts = [_text_of(p, '{%s}t' % NS['a']).strip()
                     for p in root.iter('{%s}p' % NS['a'])]
            texts = [t for t in texts if t]
            if texts:
                num = re.search(r'(\d+)', name.split('/')[-1]).group(1)
                out.append('--- 第%s页 ---' % num)
                out.extend(texts)
                total += sum(len(t) for t in texts)
            if total > max_chars:
                out.append('...[截断]')
                break
    return '\n'.join(out)


def read_xlsx(path, max_chars, max_rows=30, max_cols=12):
    """读取 xlsx：列出工作表名，并预览每表左上角若干行。"""
    out = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        shared = []
        if 'xl/sharedStrings.xml' in names:
            with z.open('xl/sharedStrings.xml') as f:
                root = ET.parse(f).getroot()
            for si in root:
                shared.append(_text_of(si, '{%s}t' % NS['ss']))
        # 工作表名
        wb = 'xl/workbook.xml'
        sheet_names = []
        if wb in names:
            with z.open(wb) as f:
                root = ET.parse(f).getroot()
            sheet_names = [s.get('name') for s in root.iter(
                '{%s}sheet' % NS['ss'])]
        out.append('工作表: ' + ' | '.join(sheet_names))
        sheets = sorted(n for n in names if re.match(r'xl/worksheets/sheet\d+\.xml$', n))
        for name in sheets[:4]:
            with z.open(name) as f:
                root = ET.parse(f).getroot()
            out.append('--- %s ---' % name)
            rows_done = 0
            for row in root.iter('{%s}row' % NS['ss']):
                cells = []
                for c in row.iter('{%s}c' % NS['ss']):
                    v = c.find('{%s}v' % NS['ss'])
                    t = c.get('t')
                    if v is None or v.text is None:
                        cells.append('')
                    elif t == 's':
                        idx = int(v.text)
                        cells.append(shared[idx] if idx < len(shared) else '')
                    else:
                        cells.append(v.text)
                cells = cells[:max_cols]
                while cells and cells[-1] == '':
                    cells.pop()
                if cells:
                    out.append(' | '.join(cells))
                rows_done += 1
                if rows_done >= max_rows:
                    out.append('...[仅预览前%d行]' % max_rows)
                    break
            if sum(len(x) for x in out) > max_chars:
                break
    return '\n'.join(out)




# ==================================================================
#  2. 公告合并（原 合并公告汇总.py）
# ==================================================================
RE_FORMAL_DAY = re.compile(r"^(.+?)（(\d{4})\.(\d{2})\.(\d{2})）(.*)$")
RE_FORMAL_ENTRY = re.compile(r"^【(.+?)】(.+?)[：:](.*)$")
RE_BX_DAY = re.compile(r"^##\s+(\d{4})-(\d{2})-(\d{2})\s*$")
RE_BX_ENTRY = re.compile(r"^-\s+\*\*(.+?)\*\*\s*[：:]\s*(.*)$")
RE_BX_HEAD = re.compile(r"^【(.+?)】(.+)$")


class MergeError(Exception):
    """工具层拒绝执行。"""



def assert_writable(out):
    if _under(out, HISTORY_DIR):
        raise MergeError("拒绝写入历史实习目录（只读）：%s" % out)
    if not any(_under(out, r) for r in WRITABLE_ROOTS):
        raise MergeError("输出必须位于 04_工作区 或 05_输出 之内：%s" % out)


# ------------------------------------------------------------- 解析旧体例
def parse_formal(text):
    """本周公司公告正式版.md → OrderedDict{YYYY-MM-DD: [(公司, 标题, 正文, 原始行)]}"""
    days, cur = {}, None
    for raw in text.split("\n"):
        line = raw.rstrip()
        if not line:
            continue
        if line.strip() == "【公司新闻】":
            continue
        m = RE_FORMAL_DAY.match(line)
        if m and "每日新闻" in line:
            cur = "%s-%s-%s" % (m.group(2), m.group(3), m.group(4))
            days.setdefault(cur, [])
            days[cur].append(("__HEAD__", "", "", line))
            continue
        m = RE_FORMAL_ENTRY.match(line)
        if m and cur:
            days[cur].append((m.group(1), m.group(2).strip(), m.group(3).strip(), line))
    return days


def parse_backup(text):
    """本周公司公告备选.md → OrderedDict{YYYY-MM-DD: [(公司, 标题, 正文, 原始行)]}"""
    days, cur = {}, None
    for raw in text.split("\n"):
        line = raw.rstrip()
        m = RE_BX_DAY.match(line)
        if m:
            cur = "%s-%s-%s" % (m.group(1), m.group(2), m.group(3))
            days.setdefault(cur, [])
            continue
        if line.startswith("## "):
            cur = None
            continue
        m = RE_BX_ENTRY.match(line)
        if not m or not cur:
            continue
        head = m.group(1).replace("（已上提正式版）", "").strip()
        hm = RE_BX_HEAD.match(head)
        if not hm:
            continue
        days[cur].append((hm.group(1), hm.group(2).strip(), m.group(2).strip(), line))
    return days


def key_of(company, title):
    """去重键：公司简称 + 公告标题（逐字，不去前缀，避免误合并同公司不同公告）。"""
    return re.sub(r"\s+", "", "%s|%s" % (company, title))


# ------------------------------------------------------------- 合并
def merge(formal, backup):
    days = {}
    for d in set(formal) | set(backup):
        rows = list(formal.get(d, []))
        seen = set(k for k in (key_of(c, t) for c, t, _b, _l in rows if c != "__HEAD__"))
        added, skipped = [], []
        for company, title, body, _line in backup.get(d, []):
            k = key_of(company, title)
            if k in seen:
                skipped.append((company, title))
                continue
            seen.add(k)
            rows.append((company, title, body, ""))
            added.append((company, title))
        days[d] = {"rows": rows, "added": added, "skipped": skipped}
    return days


def render(days, header_fmt):
    """输出汇总 md：沿用正式版的日标题体例，Word 渲染脚本可直接吃。"""
    out = []
    for d in sorted(days):
        head = None
        for c, t, _b, ln in days[d]["rows"]:
            if c == "__HEAD__":
                head = ln
                break
        if head is None:
            y, m, dd = d.split("-")
            head = header_fmt % (y, m, dd)
        out.append(head)
        out.append("【公司新闻】")
        for c, t, b, ln in days[d]["rows"]:
            if c == "__HEAD__":
                continue
            if ln and not ln.startswith("-"):
                out.append(ln)                      # 正式版原行，逐字搬运
            else:
                out.append("【%s】%s：%s" % (c, t, b))
        out.append("")
    while out and not out[-1].strip():
        out.pop()
    return "\n".join(out) + "\n"




# ==================================================================
#  3. 公告盘点（原 盘点公司公告.py）
# ==================================================================
# 公告类型关键词（按优先级从高到低匹配）
PAND_TYPE_RULES = [
    ("REITs", ["reits", "reit", "不動產投資信託", "不动产投资信托", "基础设施基金"]),
    ("盈利预警", ["盈利警告", "盈利预警", "盈利預警", "业绩预亏", "预亏", "减亏"]),
    ("业绩公告", ["業績", "业绩", "中期報告", "中期报告", "年度報告", "年报", "年報",
                  "中报", "半年报", "半年度报告", "半年度報告", "半年報",
                  "季度报告", "季報", "財務報表", "财务报表", "末期股息", "中期股息"]),
    ("销售数据", ["销售", "銷售", "签约", "簽約", "合同销售", "營運數據", "运营数据",
                  "經營數據", "经营数据", "销售情况", "销售简报", "认购",
                  "经营情况简报", "經營情況簡報", "经营简报", "主要业务经营情况",
                  "主要業務經營情況", "业务经营情况", "物業銷售"]),
    ("拿地", ["拿地", "土地", "購得", "购得", "获得房地产项目", "新增土地", "新增項目",
              "竞得", "競得", "招拍挂", "供地"]),
    ("融资", ["融资", "融資", "债券", "債券", "票据", "票據", "募集", "发行", "發行",
              "可转换", "可轉換", "贷款", "貸款", "担保", "擔保", "增发", "配股",
              "中期票据", "公司债", "超短期", "资产支持", "拆借", "财务资助",
              "財務資助", "票面利率", "回售"]),
    ("债务", ["债务", "債務", "重组", "重組", "违约", "違約", "展期", "清盘", "清盤",
              "破产", "破產", "冻结", "凍結", "輪候", "轮候", "质押", "質押", "接管"]),
    ("回购", ["回购", "回購", "增持", "减持", "減持", "股份購回"]),
    ("分红", ["分红", "股息", "派息", "权益分派", "權益分派", "红利", "紅利", "派發"]),
    ("股权变动", ["股权", "股權", "股东", "股東", "收购", "收購", "要约", "要約",
                  "出售", "转让", "轉讓", "增资", "增資", "控制权", "划转", "劃轉",
                  "限制性股票", "股权激励", "股權激勵", "无偿划转"]),
    ("管理层变动", ["董事", "高管", "高级管理人员", "高級管理人員", "总裁", "總裁",
                    "辞职", "辭職", "辞任", "辭任",
                    "聘任", "委任", "离任", "離任", "行政總裁", "首席"]),
    ("重大资产交易", ["重大资产", "重大資產", "非常重大", "主要交易", "剥离", "剝離",
                      "资产处置", "資產處置"]),
    ("重大经营事项", ["诉讼", "訴訟", "仲裁", "处罚", "處罰", "监管", "監管", "复牌",
                      "復牌", "停牌", "暂停买卖", "暫停買賣", "問詢", "问询", "说明会",
                      "异常波动", "異常波動", "风险提示", "風險提示", "计提资产减值",
                      "計提資產減值", "会计政策变更", "會計政策變更"]),
]


def pand_classify(title):
    t = title.lower()
    for label, kws in PAND_TYPE_RULES:
        for kw in kws:
            if kw.lower() in t:
                return label
    return "其他"


# A股文件名：2026-09-04：600208--衢州发展--衢州发展：总裁...pdf
RE_ASHARE = re.compile(r"^(\d{4}-\d{2}-\d{2})[：:]\s*(\d{6})--(.+?)--(.*)$")




# ==================================================================
#  4. 公司池构建（原 构建公司池.py）
# ==================================================================
TODO = "待确认"

# ---------------------------------------------------------------- 公告类型
POOL_TYPE_RULES = [
    ("REITs", ["reits", "reit", "不動產投資信託", "不动产投资信托", "基础设施基金"]),
    ("盈利预警", ["盈利警告", "盈利预警", "盈利預警", "业绩预亏", "预亏", "减亏"]),
    ("业绩公告", ["業績", "业绩", "中期報告", "中期报告", "年度報告", "年报", "年報",
                  "中报", "半年报", "半年度报告", "半年度報告", "半年報",
                  "季度报告", "季報", "財務報表", "财务报表",
                  "末期股息", "全年業績", "中期股息"]),
    ("销售数据", ["销售", "銷售", "签约", "簽約", "合同销售", "營運數據", "运营数据",
                  "經營數據", "经营数据", "销售情况", "销售简报", "认购", "物業銷售",
                  "经营情况简报", "經營情況簡報", "经营简报", "主要业务经营情况",
                  "主要業務經營情況", "业务经营情况"]),
    ("拿地", ["拿地", "土地", "購得", "购得", "获得房地产项目", "新增土地", "新增項目",
              "竞得", "競得", "招拍挂", "供地"]),
    ("融资", ["融资", "融資", "债券", "債券", "票据", "票據", "募集", "发行", "發行",
              "可转换", "可轉換", "可換股", "贷款", "貸款", "担保", "擔保", "增发",
              "配股", "中期票据", "公司债", "超短期", "资产支持", "有擔保",
              "拆借", "财务资助", "財務資助", "票面利率", "回售"]),
    ("债务", ["债务", "債務", "重组", "重組", "违约", "違約", "展期", "清盘", "清盤",
              "破产", "破產", "冻结", "凍結", "輪候", "轮候", "质押", "質押", "接管"]),
    ("回购", ["回购", "回購", "增持", "减持", "減持", "股份購回"]),
    ("分红", ["分红", "股息", "派息", "权益分派", "權益分派", "红利", "紅利", "派發"]),
    ("股权变动", ["股权", "股權", "股东", "股東", "收购", "收購", "要约", "要約",
                  "出售", "转让", "轉讓", "增资", "增資", "控制权", "股權高度集中",
                  "划转", "劃轉", "限制性股票", "股权激励", "股權激勵", "无偿划转"]),
    ("管理层变动", ["董事", "高管", "高级管理人员", "高級管理人員", "总裁", "總裁",
                    "辞职", "辭職", "辞任", "辭任",
                    "聘任", "委任", "离任", "離任", "行政總裁", "首席"]),
    ("重大资产交易", ["重大资产", "重大資產", "非常重大", "主要交易", "須予披露",
                      "须予披露", "剥离", "剝離", "资产处置", "資產處置"]),
    ("重大经营事项", ["诉讼", "訴訟", "仲裁", "处罚", "處罰", "监管", "監管", "复牌",
                      "復牌", "停牌", "暂停买卖", "暫停買賣", "問詢", "问询", "说明会",
                      "异常波动", "異常波動", "风险提示", "風險提示", "计提资产减值",
                      "計提資產減值", "会计政策变更", "會計政策變更"]),
]


def pool_classify(title):
    t = (title or "").lower()
    for label, kws in POOL_TYPE_RULES:
        for kw in kws:
            if kw.lower() in t:
                return label
    return "其他"


def norm(s):
    return re.sub(r"\s+", "", s or "")


# ------------------------------------------------- 重点房企名单（来源：使用者提供的销售数据跟踪表）
# 字段：名称, 属性分组, 表格备注
WATCHLIST = [
    ("保利发展", "央企", ""),
    ("中海集团", "央企", ""),
    ("华润置地", "央企", ""),
    ("招商蛇口", "央企", ""),
    ("建发股份", "地方国企", ""),
    ("建发国际集团", "地方国企", ""),
    ("联发集团", "地方国企", ""),
    ("绿城中国", "混合所有制", ""),
    ("绿城中国（操盘）", "混合所有制", "同一主体不同口径"),
    ("绿城中国（代建）", "混合所有制", "同一主体不同口径"),
    ("中国金茂", "央企", ""),
    ("越秀地产", "地方国企", ""),
    ("滨江集团", "民企", ""),
    ("万科A", "混合所有制", ""),
    ("保利置业集团", "央企", ""),
    ("华发股份", "地方国企", ""),
    ("龙湖集团", "民企", ""),
    ("大华集团", "民企", ""),
    ("金地集团", "混合所有制", ""),
    ("美的置业", "民企", ""),
    ("信达地产", "央企", ""),
    ("新城控股", "民企", ""),
    ("卓越集团", "民企", ""),
    ("伟星房产", "民企", ""),
    ("首开股份", "地方国企", ""),
    ("仁恒置地", "民企", ""),
    ("大悦城", "央企", ""),
    ("城建发展", "地方国企", ""),
    ("中国铁建", "建筑系", ""),
    ("中建东孚八局", "建筑系", ""),
    ("中建壹品三局", "建筑系", ""),
    ("中国中铁", "建筑系", ""),
    ("中交房地产", "建筑系", ""),
    ("中建智地一局", "建筑系", ""),
    ("电建地产", "建筑系", ""),
    ("碧桂园", "违约房企", ""),
    ("绿地控股", "违约房企", ""),
    ("世茂集团", "违约房企", ""),
    ("融创中国", "违约房企", ""),
    ("远洋集团", "违约房企", ""),
    ("中国恒大", "违约房企", ""),
    ("旭辉控股集团", "违约房企", ""),
    ("宝龙地产", "违约房企", ""),
]

# ------------------------------------------------- 重点房企 ↔ 本地历史公告 匹配表
# 只登记「已在本工作区提取结果中实际出现」的公司；未登记者一律待确认。
# A股：公告文件名中的证券简称完全相同（代码唯一，可作键）
WATCH_MATCH = {
    "保利发展": ("A股", "600048"),
    "万科A": ("A股", "000002"),
    "招商蛇口": ("A股", "001979"),
    "华发股份": ("A股", "600325"),
    "滨江集团": ("A股", "002244"),
    "金地集团": ("A股", "600383"),
    "信达地产": ("A股", "600657"),
    "新城控股": ("A股", "601155"),
    "首开股份": ("A股", "600376"),
    "大悦城": ("A股", "000031"),
    "城建发展": ("A股", "600266"),
    "绿地控股": ("A股", "600606"),
}

# 港股：用公司全称关键字匹配。不用代码作键 —— 部分港股未提取到代码，
# 若以空代码作键会与所有「无代码」公司撞车（曾导致误标）。
WATCH_HK_KEYWORDS = {
    "保利置業集團": "保利置业集团",
    "華潤置地": "华润置地",
    "中國金茂": "中国金茂",
    "越秀地產": "越秀地产",
    "龍湖集團": "龙湖集团",
    "美的置業": "美的置业",
    "碧桂園控股": "碧桂园",          # 不含「碧桂園服務控股」（物业平台）
    "旭輝控股": "旭辉控股集团",
    "寶龍地產": "宝龙地产",
    "綠城中國": "绿城中国",
}


def watch_hk(name):
    """按关键字判断港股公司是否属于重点名单，返回重点名单中的名称或 ''。"""
    for kw, wname in WATCH_HK_KEYWORDS.items():
        if kw in name:
            return wname
    return ""

# 重点房企中，本地只找到「同集团物业/服务上市平台」而非母公司本身，需人工确认
WATCH_RELATED = {
    "融创中国": "本地仅见「融創服務控股有限公司」(01516)，为物业平台，非母公司",
    "远洋集团": "本地仅见「遠洋服務控股有限公司」(06677)，为物业平台，非母公司",
    "中国恒大": "本地仅见「恒大物業集團有限公司」(06666)，为物业平台，非母公司",
    "中海集团": "本地见「中國海外宏洋集團有限公司」并附代号00688，"
                "公司名与代号是否对应需人工复核",
}


def load_ashare():
    p = os.path.join(PAND, "A股公司统计.csv")
    out = {}
    for r in csv.DictReader(open(p, encoding="utf-8-sig")):
        out[r["代码"]] = {"名称": r["简称"], "次数": int(r["公告数"]),
                          "类型": r["主要类型"], "样例": r["样例"]}
    return out


def load_hk():
    """返回 公司名 -> {codes:set, count:int, types:Counter, samples:list}"""
    p = os.path.join(PAND, "港股公司提取.csv")
    recs = list(csv.DictReader(open(p, encoding="utf-8-sig")))

    # 代号 -> 已可靠识别出的公司名（同源互证，用于补齐缺名记录）
    code2name = defaultdict(set)
    for r in recs:
        if r["公司名称"] and r["股份代号"]:
            code2name[r["股份代号"]].add(r["公司名称"])
    code2name = {k: next(iter(v)) for k, v in code2name.items() if len(v) == 1}

    agg = defaultdict(lambda: {"codes": set(), "count": 0,
                               "types": defaultdict(int), "samples": []})
    unresolved = []
    for r in recs:
        name = r["公司名称"]
        via = ""
        if not name and r["股份代号"] in code2name:
            name = code2name[r["股份代号"]]
            via = "（公司名由同源公告的股份代号关联，待人工复核）"
        if not name:
            unresolved.append(r)
            continue
        e = agg[name]
        if r["股份代号"]:
            e["codes"].add(r["股份代号"])
        e["count"] += 1
        e["types"][pool_classify(r["标题"])] += 1
        if len(e["samples"]) < 3:
            e["samples"].append("%s/%s/%s" % (r["周"], r["日期"], r["标题"][:50]))
        if via:
            e.setdefault("notes", set()).add(via)
    return agg, unresolved




# ==================================================================
#  5. 港股公司提取（原 提取港股公司.py）
# ==================================================================
CJK = "\u4e00-\u9fff"
FW = "\uff08\uff09\uff3b\uff3d"
DOT = "\u00b7\uff0e\uff65"

# 必须整段剔除的样板/第三方机构片段
EXCLUDE_SUBSTR = [
    "香港交易及結算所有限公司", "香港聯合交易所有限公司",
    "香港交易及结算所有限公司", "香港联合交易所有限公司",
    "香港中央結算有限公司", "香港中央结算有限公司",
    "香港中央證券登記有限公司", "香港中央证券登记有限公司",
    "卓佳證券登記有限公司", "卓佳证券登记有限公司",
    "註冊成立", "注册成立",
    "會計師事務所", "会计师事务所", "律師事務所", "律师事务所",
]

# 名称中不得出现的字样（避免把报告标题/第三方机构当成公司名）
NAME_FORBIDDEN = [
    "本公司", "本集團", "本集团", "附屬", "附属",
    "公告", "報告", "报告", "目錄", "目录", "主席", "年報", "年报", "中期",
    "季度", "摘要", "聲明", "声明", "職能", "职能", "名單", "名单",
    "會計師", "会计师", "律師", "律师", "證券登記", "证券登记", "代理人",
    "環境衛生", "环境卫生", "電影", "电影", "製片", "制片", "銀行", "银行",
]

RE_CODE = re.compile(r"股份代號\s*[：:﹕]\s*(\d{3,5})(?!\d)")
RE_CODE_LOOSE = re.compile(r"股份代號[^0-9]{0,6}(\d{3,5})(?!\d)")
RE_CODE_FIELD = re.compile(r"證券?代號[^0-9]{0,12}(\d{3,5})(?!\d)")
RE_CODE_EN = re.compile(r"Stock\s*Code[^0-9]{0,6}(\d{3,5})(?!\d)", re.I)
# 整行就是中文公司全称
RE_NAME_LINE = re.compile(r"^[%s%s%s]{4,30}(?:有限公司|股份有限公司)$" % (CJK, FW, DOT))


def hk_clean(s):
    return re.sub(r"[\s\u3000]+", "", s or "")


def bad_name(name):
    if len(name) < 4 or len(name) > 30:
        return True
    for bad in EXCLUDE_SUBSTR + NAME_FORBIDDEN:
        if bad in name:
            return True
    return False


def extract(path):
    import pymupdf
    doc = pymupdf.open(path)
    try:
        lines = []
        for i in range(min(2, doc.page_count)):
            lines.extend(doc[i].get_text().splitlines())
    finally:
        doc.close()

    raw = "\n".join(lines)

    # --- 股份代号：在原始文本上匹配，避免数字与后文粘连 ---
    code = ""
    m = (RE_CODE.search(raw) or RE_CODE_LOOSE.search(raw)
         or RE_CODE_FIELD.search(raw) or RE_CODE_EN.search(raw))
    if m:
        code = m.group(1).zfill(5) if len(m.group(1)) < 5 else m.group(1)

    # --- 公司名规则 A：月報表结构化字段「公司名稱:」 ---
    name = ""
    for i, line in enumerate(lines):
        c = hk_clean(line)
        if c in ("公司名稱:", "公司名稱：", "公司名称:", "公司名称："):
            for j in range(i + 1, min(i + 4, len(lines))):
                cand = hk_clean(lines[j])
                if not cand:
                    continue
                cand = re.sub(r"[（(].*$", "", cand).strip()
                if RE_NAME_LINE.match(cand) and not bad_name(cand):
                    name = cand
                break
            if name:
                break

    # --- 公司名规则 B：整行即公司全称 ---
    if name:
        return code, name
    for line in lines:
        c = hk_clean(line)
        if not c or not RE_NAME_LINE.match(c):
            continue
        if bad_name(c):
            continue
        name = c
        break
    return code, name




# ==================================================================
#  6. 本周正文抽取（原 提取本周文本.py）
# ==================================================================
PREVIEW_INDEX = 300     # _index.tsv 里的预览长度
PREVIEW_BODY = 4000     # bodies.txt 里的每文件正文上限

RE_A_NAME = re.compile(r"^(\d{4}-\d{2}-\d{2})[：:]\s*(\d{6})--(.+?)--(.*)$")
RE_A_NAME_NOCODE = re.compile(r"^(\d{4}-\d{2}-\d{2})[：:]\s*(.*)$")


def week_clean(s):
    return re.sub(r"[ \t　]+", " ", (s or "").replace("\r", "")).strip()


def pdf_text(path, max_chars=None, max_pages=None):
    """只读提取 PDF 文本。返回 (完整截断文本, 页数)。"""
    import pymupdf
    doc = pymupdf.open(path)
    try:
        n = doc.page_count
        pages = n if max_pages is None else min(max_pages, n)
        parts = []
        total = 0
        for i in range(pages):
            t = doc[i].get_text()
            parts.append(t)
            total += len(t)
            if max_chars and total >= max_chars:
                break
        text = "\n".join(parts)
    finally:
        doc.close()
    if max_chars and len(text) > max_chars:
        text = text[:max_chars] + "\n...[截断]"
    return text, n


def docx_text(path, max_chars=None, max_pages=None):
    import zipfile
    from xml.etree import ElementTree as ET
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as z:
        with z.open("word/document.xml") as f:
            root = ET.parse(f).getroot()
    lines = []
    for p in root.iter(W + "p"):
        line = week_clean("".join(t.text or "" for t in p.iter(W + "t")))
        if line:
            lines.append(line)
    text = "\n".join(lines)
    if max_chars and len(text) > max_chars:
        text = text[:max_chars] + "\n...[截断]"
    return text, 1


def read_any(path, max_chars=None, max_pages=None):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return pdf_text(path, max_chars, max_pages)
    if ext in (".docx", ".doc"):
        return docx_text(path, max_chars, max_pages)
    if ext in (".txt", ".md"):
        with open(path, encoding="utf-8", errors="replace") as f:
            t = f.read()
        return (t[:max_chars] if max_chars else t), 1
    return "", 0


def walk(root):
    out = []
    for dirpath, _dirs, files in os.walk(root):
        for fn in sorted(files):
            if fn.lower().endswith((".pdf", ".docx", ".doc", ".txt", ".md")):
                out.append(os.path.join(dirpath, fn))
    return sorted(out)


def hk_company(path):
    """复用本模块的港股公司名/代号规则（只读前 2 页）。"""
    try:
        return extract(path)
    except Exception:
        return "", ""




# ==================================================================
#  7. Excel 工具（原 excel工具.py）
# ==================================================================
class ExcelToolError(Exception):
    """工具层拒绝执行（安全约束或调用错误）。"""



def assert_readable(src):
    """源文件可读即可（历史实习目录允许**读**，只是禁止写）。"""
    if not os.path.isfile(src):
        raise ExcelToolError("源文件不存在：%s" % src)
    if not src.lower().endswith((".xlsx", ".xlsm")):
        raise ExcelToolError("只支持 .xlsx / .xlsm：%s" % src)


def assert_writable_target(out, src=None):
    """写入目标必须落在工作区内的可写目录，且不等于源文件。"""
    if not out:
        raise ExcelToolError("必须显式指定输出路径 --out（默认不覆盖原文件）")

    if _under(out, HISTORY_DIR):
        raise ExcelToolError(
            "拒绝写入历史实习目录（只读）：%s" % out)

    if not any(_under(out, r) for r in WRITABLE_ROOTS):
        raise ExcelToolError(
            "输出路径必须位于 04_工作区 或 05_输出 之内：%s\n"
            "  允许的根目录：\n    - %s"
            % (out, "\n    - ".join(WRITABLE_ROOTS)))

    if src and _abs(out).lower() == _abs(src).lower():
        raise ExcelToolError("拒绝覆盖源文件：%s" % out)

    if not out.lower().endswith((".xlsx", ".xlsm")):
        raise ExcelToolError("输出文件必须是 .xlsx / .xlsm：%s" % out)


# ----------------------------------------------------------- 基础操作
def open_book(path, data_only=False):
    """打开工作簿。

    data_only=False（默认）保留公式；True 读取公式的**缓存值**（无缓存则为 None）。
    """
    assert_readable(path)
    try:
        return load_workbook(filename=path, data_only=data_only)
    except Exception as exc:                       # noqa: BLE001
        raise ExcelToolError("打开失败：%s（%s）" % (path, exc))


def list_sheets(wb):
    """返回 sheet 名称列表。"""
    return list(wb.sheetnames)


def get_sheet(wb, name=None):
    """取 sheet；name 为空则取活动表。找不到 → 报错（不静默 fallback）。"""
    if name is None or name == "":
        return wb.active
    if name not in wb.sheetnames:
        raise ExcelToolError(
            "工作表不存在：%r\n  可用的有：%s" % (name, "、".join(wb.sheetnames)))
    return wb[name]


def read_cell(wb, sheet, ref, data_only=False):
    """读单个单元格。

    若已用 data_only=True 打开，传入的 wb 已带缓存值，这里不再重复加载。
    """
    ws = get_sheet(wb, sheet)
    try:
        return ws[ref].value
    except Exception as exc:                       # noqa: BLE001
        raise ExcelToolError("读取单元格失败 %s!%s（%s）" % (sheet, ref, exc))


def read_range(wb, sheet, ref):
    """读区域，返回二维列表（行优先）。单格也返回 [[值]]。"""
    ws = get_sheet(wb, sheet)
    try:
        min_c, min_r, max_c, max_r = range_boundaries(ref)
    except Exception as exc:                       # noqa: BLE001
        raise ExcelToolError("区域格式不合法：%r（%s）" % (ref, exc))
    if min_r == max_r and min_c == max_c:
        return [[ws.cell(row=min_r, column=min_c).value]]
    return [[ws.cell(row=r, column=c).value
             for c in range(min_c, max_c + 1)]
            for r in range(min_r, max_r + 1)]


def write_cell(wb, sheet, ref, value):
    """写单个单元格（只改内存中的工作簿，落盘需再调用 save_as）。"""
    ws = get_sheet(wb, sheet)
    try:
        ws[ref] = value
    except Exception as exc:                       # noqa: BLE001
        raise ExcelToolError("写入单元格失败 %s!%s（%s）" % (sheet, ref, exc))
    return ws


def save_as(wb, out, src=None, force=False):
    """另存为新文件。默认拒绝覆盖已存在的文件（除非 force=True）。"""
    assert_writable_target(out, src=src)
    if os.path.exists(out) and not force:
        raise ExcelToolError(
            "目标文件已存在，拒绝覆盖：%s\n  如需覆盖请显式加 --force" % out)
    os.makedirs(os.path.dirname(_abs(out)), exist_ok=True)
    try:
        wb.save(out)
    except Exception as exc:                       # noqa: BLE001
        raise ExcelToolError("保存失败：%s（%s）" % (out, exc))
    return out


def copy_to(src, out, force=False):
    """把源工作簿另存为副本。源文件只读，绝不被改动。"""
    wb = open_book(src)
    return save_as(wb, out, src=src, force=force)


# ----------------------------------------------------------- 命令行
def _coerce(text, kind):
    if kind == "str" or text is None:
        return text
    if kind == "int":
        return int(text)
    if kind == "float":
        return float(text)
    # auto：能转成数字就转，否则原样字符串
    try:
        return int(text)
    except (TypeError, ValueError):
        pass
    try:
        return float(text)
    except (TypeError, ValueError):
        return text


def _fmt(v):
    if v is None:
        return "(空)"
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, dt.date):
        return v.isoformat()
    return str(v)



def _warn_lossy():
    log("")
    log("[注意] openpyxl 重新保存会丢失：图表、数据透视表、宏(VBA)、"
        "部分条件格式与图片。")
    log("       若源文件含上述内容，请勿用本工具处理，改由人工另存。")




# ==================================================================
#  8. 依赖检查（原 检查依赖.py）
# ==================================================================
DEPS = [
    ("python-docx", "docx",   "读取 Word（公告模板、周报）"),
    ("python-pptx", "pptx",   "生成 PPT 页"),
    ("openpyxl",    "openpyxl", "Excel 结构化数据更新"),
    ("lxml",        "lxml",   "XML 解析支撑"),
    ("Pillow",      "PIL",    "图片处理支撑"),
]


def check_dependencies():
    missing = []
    for display, mod, _purpose in DEPS:
        try:
            m = importlib.import_module(mod)
            ver = getattr(m, "__version__", "?")
            print("OK|%s|%s" % (display, ver))
        except ImportError:
            print("MISS|%s|" % display)
            missing.append(display)

    if missing:
        print("MISSING_TOTAL|%d" % len(missing))
        return 1
    return 0
